"""GSTR-1 sections, GSTR-3B summary, the reconciliation between them, and reports."""
import csv, io
from decimal import Decimal
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from .. import models as M
from ..deps import Ctx, need
from ..service import invoice_tax, vinv_tax, settled, live_invoices
from ..tax import q2

router = APIRouter(prefix="/returns", tags=["returns"])
MON = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]


def gd(d):
    return f"{d.day:02d}-{MON[d.month - 1]}-{d.year}" if d else ""


def _pos(ctx, code):
    s = ctx.db.get(M.State, code)
    return f"{code}-{s.name}" if s else code


def _tax_invoices(ctx, period):
    q = live_invoices(ctx).order_by(M.Invoice.doc_date, M.Invoice.id)
    return [i for i in ctx.db.execute(q).scalars() if str(i.doc_date)[:7] == period]


def _cancelled(ctx, period):
    q = ctx.scope(select(M.Invoice), M.Invoice).where(M.Invoice.doc_type == "TAX",
                                                       M.Invoice.status == "CANCELLED")
    return [i for i in ctx.db.execute(q).scalars() if str(i.doc_date)[:7] == period]


@router.get("/periods")
def periods(ctx: Ctx = Depends(need("gstr"))):
    p = {str(i.doc_date)[:7] for i in ctx.db.execute(ctx.scope(select(M.Invoice), M.Invoice)).scalars()}
    p |= {str(v.doc_date)[:7] for v in ctx.db.execute(ctx.scope(select(M.VendorInvoice), M.VendorInvoice)).scalars()}
    return sorted(p, reverse=True)


@router.get("/gstr1")
def gstr1(period: str = Query(..., description="YYYY-MM"), ctx: Ctx = Depends(need("gstr"))):
    """The return as it would be filed, split into its sections."""
    sec = {"b2b": [], "b2cl": [], "b2cs": []}
    tot = {"taxable": Decimal(0), "cgst": Decimal(0), "sgst": Decimal(0), "igst": Decimal(0)}
    for inv in _tax_invoices(ctx, period):
        t = invoice_tax(ctx, inv)
        for k in tot:
            tot[k] += getattr(t, k)
        row = {"doc_no": inv.doc_no, "doc_date": gd(inv.doc_date),
               "recipient": inv.customer.name, "gstin": inv.gstin,
               "pos": _pos(ctx, inv.pos_state), "invoice_value": float(t.rounded),
               "taxable": float(t.taxable), "igst": float(t.igst),
               "cgst": float(t.cgst), "sgst": float(t.sgst),
               "reverse_chg": inv.reverse_chg}
        if inv.gstin:
            sec["b2b"].append(row)
        elif not t.intra and t.rounded > 250000:
            sec["b2cl"].append(row)
        else:
            sec["b2cs"].append(row)
    proforma = len([i for i in ctx.db.execute(ctx.scope(select(M.Invoice), M.Invoice).where(M.Invoice.doc_type == "PRO")).scalars()
        if str(i.doc_date)[:7] == period])
    cancelled = _cancelled(ctx, period)
    return {"period": period, "sections": sec,
            "totals": {k: float(q2(v)) for k, v in tot.items()},
            "proforma_excluded": proforma,
            "cancelled": [{"doc_no": i.doc_no, "doc_date": gd(i.doc_date),
                           "reason": i.cancel_reason} for i in cancelled]}


@router.get("/gstr3b")
def gstr3b(period: str = Query(...), ctx: Ctx = Depends(need("gstr"))):
    out = {"taxable": Decimal(0), "cgst": Decimal(0), "sgst": Decimal(0), "igst": Decimal(0)}
    n_out = 0
    for inv in _tax_invoices(ctx, period):
        t = invoice_tax(ctx, inv)
        n_out += 1
        for k in out:
            out[k] += getattr(t, k)
    inp = {"taxable": Decimal(0), "cgst": Decimal(0), "sgst": Decimal(0), "igst": Decimal(0)}
    n_in = 0
    vinvs = [v for v in ctx.db.execute(ctx.scope(select(M.VendorInvoice), M.VendorInvoice)).scalars()
             if str(v.doc_date)[:7] == period]
    for vi in vinvs:
        t = vinv_tax(ctx, vi)
        n_in += 1
        for k in inp:
            inp[k] += getattr(t, k)
    net = {k: max(Decimal(0), out[k] - inp[k]) for k in ("cgst", "sgst", "igst")}
    return {"period": period,
            "outward": {**{k: float(q2(v)) for k, v in out.items()}, "documents": n_out},
            "itc": {**{k: float(q2(v)) for k, v in inp.items()}, "documents": n_in,
                    "source": "vendor invoices"},
            "net_payable": {k: float(q2(v)) for k, v in net.items()},
            "caveat": ("Input credit is taken from vendor invoices recorded here, not from "
                       "GSTR-2B. Blocked credits under section 17(5), reverse charge and "
                       "credit reversals are not handled.")}


@router.get("/reconciliation")
def reconciliation(period: str = Query(...), ctx: Ctx = Depends(need("gstr"))):
    g1 = gstr1(period, ctx)
    g3 = gstr3b(period, ctx)
    rows = []
    for head in ("taxable", "igst", "cgst", "sgst"):
        a, b = g1["totals"][head], g3["outward"][head]
        rows.append({"head": head, "gstr1": a, "gstr3b": b,
                     "difference": round(a - b, 2), "agrees": abs(a - b) < 0.5})
    return {"period": period, "rows": rows, "all_agree": all(r["agrees"] for r in rows),
            "note": ("Both are built from the same set of tax invoices, so a difference "
                     "here indicates a data problem rather than a reconciling item. Credit "
                     "and debit notes, advances, exports and amendments are not covered.")}


def _csv(rows, filename):
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/gstr1/{section}.csv")
def gstr1_csv(section: str, period: str = Query(...), ctx: Ctx = Depends(need("gstr"))):
    """Section CSVs in the layout the GST Returns Offline Tool imports."""
    g = gstr1(period, ctx)
    if section == "b2b":
        rows = [["GSTIN/UIN of Recipient", "Receiver Name", "Invoice Number", "Invoice date",
                 "Invoice Value", "Place Of Supply", "Reverse Charge",
                 "Applicable % of Tax Rate", "Invoice Type", "E-Commerce GSTIN",
                 "Rate", "Taxable Value", "Cess Amount"]]
        for r in g["sections"]["b2b"]:
            rate = 18 if r["taxable"] else 0
            rows.append([r["gstin"], r["recipient"], r["doc_no"], r["doc_date"],
                         r["invoice_value"], r["pos"], r["reverse_chg"], "", "Regular B2B",
                         "", rate, r["taxable"], ""])
    elif section == "b2cs":
        rows = [["Type", "Place Of Supply", "Rate", "Applicable % of Tax Rate",
                 "Taxable Value", "Cess Amount", "E-Commerce GSTIN"]]
        agg = {}
        for r in g["sections"]["b2cs"]:
            agg[r["pos"]] = agg.get(r["pos"], 0) + r["taxable"]
        for pos, val in agg.items():
            rows.append(["OE", pos, 18, "", round(val, 2), "", ""])
    elif section == "hsn":
        rows = [["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Rate",
                 "Taxable Value", "Integrated Tax Amount", "Central Tax Amount",
                 "State/UT Tax Amount", "Cess Amount"]]
        agg = {}
        for inv in _tax_invoices(ctx, period):
            t = invoice_tax(ctx, inv)
            for L in t.lines:
                k = (L.hsn, str(L.rate), L.uom)
                a = agg.setdefault(k, {"d": L.descr, "q": Decimal(0), "v": Decimal(0),
                                       "t": Decimal(0), "i": Decimal(0), "c": Decimal(0),
                                       "s": Decimal(0)})
                a["q"] += L.qty
                a["t"] += L.amount
                a["v"] += L.amount + L.cgst + L.sgst + L.igst
                a["i"] += L.igst
                a["c"] += L.cgst
                a["s"] += L.sgst
        for (hsn, rate, uom), a in agg.items():
            rows.append([hsn, a["d"], uom, float(a["q"]), float(q2(a["v"])), rate,
                         float(q2(a["t"])), float(q2(a["i"])), float(q2(a["c"])),
                         float(q2(a["s"])), ""])
    else:
        nos = sorted([r["doc_no"] for s in g["sections"].values() for r in s]
                     + [c["doc_no"] for c in g["cancelled"]])
        rows = [["Nature of Document", "Sr. No. From", "Sr. No. To", "Total Number", "Cancelled"],
                ["Invoices for outward supply", nos[0] if nos else "",
                 nos[-1] if nos else "", len(nos), len(g["cancelled"])]]
    return _csv(rows, f"{section}_{period}.csv")


# ------------------------------------------------------------- reports
@router.get("/reports/receivables")
def receivables(ctx: Ctx = Depends(need("reports"))):
    out = []
    for inv in ctx.db.execute(live_invoices(ctx)).scalars():
        t = invoice_tax(ctx, inv)
        rec = settled(ctx, invoice_id=inv.id)
        bal = t.rounded - rec
        if bal <= Decimal("0.5"):
            continue
        out.append({"customer": inv.customer.name, "doc_no": inv.doc_no,
                    "doc_date": inv.doc_date, "due_date": inv.due_date,
                    "total": float(t.rounded), "received": float(rec),
                    "outstanding": float(bal)})
    return sorted(out, key=lambda r: r["doc_date"])


@router.get("/reports/payables")
def payables(ctx: Ctx = Depends(need("gstr"))):
    out = []
    for vi in ctx.db.execute(ctx.scope(select(M.VendorInvoice), M.VendorInvoice)).scalars():
        t = vinv_tax(ctx, vi)
        paid = settled(ctx, vinv_id=vi.id)
        bal = t.rounded - paid
        if bal <= Decimal("0.5"):
            continue
        out.append({"vendor": vi.vendor.name, "doc_no": vi.doc_no, "doc_date": vi.doc_date,
                    "due_date": vi.due_date, "total": float(t.rounded),
                    "paid": float(paid), "outstanding": float(bal)})
    return sorted(out, key=lambda r: r["doc_date"])


@router.get("/reports/margin")
def margin(ctx: Ctx = Depends(need("gstr"))):
    out = []
    for m in ctx.db.execute(ctx.scope(select(M.Material), M.Material).order_by(M.Material.code)).scalars():
        sq = sv = bq = bv = Decimal(0)
        for l in ctx.db.execute(select(M.InvoiceLine).where(M.InvoiceLine.material_id == m.id)).scalars():
            if l.invoice.doc_type == "TAX" and l.invoice.status != "CANCELLED":
                sq += l.qty
                sv += l.qty * l.price
        for l in ctx.db.execute(select(M.VendorInvoiceLine).where(M.VendorInvoiceLine.material_id == m.id)).scalars():
            bq += l.qty
            bv += l.qty * l.price
        if not sq and not bq:
            continue
        out.append({"code": m.code, "descr": m.descr, "sold_qty": float(sq),
                    "sales_value": float(q2(sv)), "bought_qty": float(bq),
                    "purchase_value": float(q2(bv)), "margin": float(q2(sv - bv)),
                    "margin_pct": float(q2((sv - bv) / sv * 100)) if sv else None})
    return out


# =========================================== GSTR-1 for the GST portal
def _rate_split(t):
    """Per-rate totals of one document: {rate: (taxable, igst, cgst, sgst)}."""
    agg = {}
    for L in t.lines:
        a = agg.setdefault(L.rate, [Decimal(0)] * 4)
        a[0] += L.amount; a[1] += L.igst; a[2] += L.cgst; a[3] += L.sgst
    return {r: tuple(q2(x) for x in v) for r, v in agg.items()}


def _fp(period):
    y, m = period.split("-")
    return f"{m}{y}"


def _portal_date(d):
    return d.strftime("%d-%m-%Y")


def gstr1_portal(ctx, period):
    """GSTR-1 in the JSON schema the GST portal / Returns Offline Tool accepts:
    b2b, b2cl, b2cs, hsn and doc_issue. Cancelled invoices are counted in
    doc_issue and nowhere else."""
    org = ctx.tenant
    invs = _tax_invoices(ctx, period)
    b2b, b2cl, b2cs, hsn = {}, {}, {}, {}
    for inv in invs:
        t = invoice_tax(ctx, inv)
        split = _rate_split(t)
        intra = t.intra
        items = [{"num": i + 1, "itm_det": {"rt": float(r), "txval": float(v[0]),
                  **({"camt": float(v[2]), "samt": float(v[3])} if intra else {"iamt": float(v[1])}),
                  "csamt": 0}} for i, (r, v) in enumerate(sorted(split.items()))]
        rec = {"inum": inv.doc_no, "idt": _portal_date(inv.doc_date), "val": float(t.rounded),
               "pos": inv.pos_state, "rchrg": inv.reverse_chg, "inv_typ": "R", "itms": items}
        if inv.gstin:
            b2b.setdefault(inv.gstin, []).append(rec)
        elif not intra and t.rounded > 250000:
            b2cl.setdefault(inv.pos_state, []).append(rec)
        else:
            for r, v in split.items():
                k = (inv.pos_state, r)
                a = b2cs.setdefault(k, {"sply_ty": "INTRA" if intra else "INTER", "pos": inv.pos_state,
                                        "typ": "OE", "rt": float(r), "txval": 0.0, "iamt": 0.0,
                                        "camt": 0.0, "samt": 0.0, "csamt": 0})
                a["txval"] = float(q2(Decimal(str(a["txval"])) + v[0]))
                a["iamt"] = float(q2(Decimal(str(a["iamt"])) + v[1]))
                a["camt"] = float(q2(Decimal(str(a["camt"])) + v[2]))
                a["samt"] = float(q2(Decimal(str(a["samt"])) + v[3]))
        for L in t.lines:
            k = (L.hsn, L.rate, L.uom)
            h = hsn.setdefault(k, {"hsn_sc": L.hsn, "desc": L.descr[:30], "uqc": _uqc(L.uom),
                                   "qty": Decimal(0), "txval": Decimal(0), "iamt": Decimal(0),
                                   "camt": Decimal(0), "samt": Decimal(0), "rt": float(L.rate)})
            h["qty"] += L.qty; h["txval"] += L.amount; h["iamt"] += L.igst
            h["camt"] += L.cgst; h["samt"] += L.sgst
    cancelled = _cancelled(ctx, period)
    all_nos = sorted([i.doc_no for i in invs] + [i.doc_no for i in cancelled])
    doc = {"gstin": org.gstin or "", "fp": _fp(period), "version": "GST3.0.4", "hash": "hash",
           "b2b": [{"ctin": g, "inv": v} for g, v in b2b.items()],
           "b2cl": [{"pos": p, "inv": v} for p, v in b2cl.items()],
           "b2cs": list(b2cs.values()),
           "hsn": {"data": [{"num": i + 1, **{k: (float(q2(v)) if isinstance(v, Decimal) else v)
                                              for k, v in h.items()}, "csamt": 0}
                            for i, h in enumerate(hsn.values())]},
           "doc_issue": {"doc_det": [{"doc_num": 1, "docs": [{"num": 1,
                         "from": all_nos[0] if all_nos else "", "to": all_nos[-1] if all_nos else "",
                         "totnum": len(all_nos), "cancel": len(cancelled),
                         "net_issue": len(all_nos) - len(cancelled)}]}]}}
    return doc


UQC = {"NOS": "NOS", "PCS": "PCS", "BOX": "BOX", "KG": "KGS", "KGS": "KGS", "LTR": "LTR",
       "MTR": "MTR", "SET": "SET", "PAC": "PAC", "PAIR": "PRS", "EA": "NOS", "UNT": "UNT"}


def _uqc(uom):
    return UQC.get((uom or "").upper(), "OTH")


@router.get("/gstr1/portal.json")
def gstr1_json(period: str = Query(...), ctx: Ctx = Depends(need("gstr"))):
    import json
    body = json.dumps(gstr1_portal(ctx, period), indent=1)
    return StreamingResponse(iter([body]), media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="GSTR1_{_fp(period)}.json"'})


@router.get("/gstr1/portal.xlsx")
def gstr1_xlsx(period: str = Query(...), ctx: Ctx = Depends(need("gstr"))):
    """One workbook, one sheet per GSTR-1 table, headings as in the Returns
    Offline Tool template (b2b, b2cl, b2cs, hsn, docs)."""
    from openpyxl import Workbook
    d = gstr1_portal(ctx, period)
    wb = Workbook()
    ws = wb.active; ws.title = "b2b"
    ws.append(["GSTIN/UIN of Recipient", "Receiver Name", "Invoice Number", "Invoice date",
               "Invoice Value", "Place Of Supply", "Reverse Charge", "Applicable % of Tax Rate",
               "Invoice Type", "E-Commerce GSTIN", "Rate", "Taxable Value", "Cess Amount"])
    names = {i.gstin: i.customer.name for i in _tax_invoices(ctx, period) if i.gstin}
    for g in d["b2b"]:
        for inv in g["inv"]:
            for it in inv["itms"]:
                ws.append([g["ctin"], names.get(g["ctin"], ""), inv["inum"], inv["idt"], inv["val"],
                           _pos(ctx, inv["pos"]), inv["rchrg"], "", "Regular B2B", "",
                           it["itm_det"]["rt"], it["itm_det"]["txval"], 0])
    ws = wb.create_sheet("b2cl")
    ws.append(["Invoice Number", "Invoice date", "Invoice Value", "Place Of Supply",
               "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"])
    for p in d["b2cl"]:
        for inv in p["inv"]:
            for it in inv["itms"]:
                ws.append([inv["inum"], inv["idt"], inv["val"], _pos(ctx, p["pos"]), "",
                           it["itm_det"]["rt"], it["itm_det"]["txval"], 0, ""])
    ws = wb.create_sheet("b2cs")
    ws.append(["Type", "Place Of Supply", "Applicable % of Tax Rate", "Rate", "Taxable Value",
               "Cess Amount", "E-Commerce GSTIN"])
    for r in d["b2cs"]:
        ws.append([r["typ"], _pos(ctx, r["pos"]), "", r["rt"], r["txval"], 0, ""])
    ws = wb.create_sheet("hsn")
    ws.append(["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Taxable Value",
               "Integrated Tax Amount", "Central Tax Amount", "State/UT Tax Amount", "Cess Amount",
               "Rate"])
    for h in d["hsn"]["data"]:
        ws.append([h["hsn_sc"], h["desc"], h["uqc"], h["qty"],
                   round(h["txval"] + h["iamt"] + h["camt"] + h["samt"], 2), h["txval"],
                   h["iamt"], h["camt"], h["samt"], 0, h["rt"]])
    ws = wb.create_sheet("docs")
    ws.append(["Nature of Document", "Sr. No. From", "Sr. No. To", "Total Number", "Cancelled"])
    dd = d["doc_issue"]["doc_det"][0]["docs"][0]
    ws.append(["Invoices for outward supply", dd["from"], dd["to"], dd["totnum"], dd["cancel"]])
    for w in wb.worksheets:
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(14, min(40,
                max(len(str(c.value or "")) for c in col) + 2))
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="GSTR1_{_fp(period)}.xlsx"'})


# ================================================= reconciliation A − B
@router.get("/reconciliation/ab")
def reconciliation_ab(period: str = Query(...), ctx: Ctx = Depends(need("gstr"))):
    """A = GSTR-1 built from the books (what will be uploaded to the portal).
    B = GSTR-3B as uploaded here (what was, or will be, filed).
    Each head shows A, B and A − B. Where no GSTR-3B has been uploaded for the
    period, B falls back to the GSTR-3B computed from the books and is flagged."""
    g1 = gstr1(period, ctx)
    filed = ctx.db.execute(ctx.scope(select(M.Gstr3bUpload), M.Gstr3bUpload)
                           .where(M.Gstr3bUpload.period == period)).scalar_one_or_none()
    g3 = gstr3b(period, ctx)
    if filed:
        b_out = {"taxable": float(filed.out_taxable), "igst": float(filed.out_igst),
                 "cgst": float(filed.out_cgst), "sgst": float(filed.out_sgst)}
        b_itc = {"igst": float(filed.itc_igst), "cgst": float(filed.itc_cgst),
                 "sgst": float(filed.itc_sgst)}
    else:
        b_out = {k: g3["outward"][k] for k in ("taxable", "igst", "cgst", "sgst")}
        b_itc = {k: g3["itc"][k] for k in ("igst", "cgst", "sgst")}
    rows = []
    for head, lbl in (("taxable", "Outward taxable value (3.1a)"), ("igst", "Outward IGST"),
                      ("cgst", "Outward CGST"), ("sgst", "Outward SGST")):
        a, b = g1["totals"][head], b_out[head]
        rows.append({"section": "Outward", "head": lbl, "a": a, "b": b,
                     "a_minus_b": round(a - b, 2), "agrees": abs(a - b) <= 1})
    for head, lbl in (("igst", "Input credit IGST (4A)"), ("cgst", "Input credit CGST"),
                      ("sgst", "Input credit SGST")):
        a, b = g3["itc"][head], b_itc[head]
        rows.append({"section": "Input credit", "head": lbl, "a": a, "b": b,
                     "a_minus_b": round(a - b, 2), "agrees": abs(a - b) <= 1})
    return {"period": period, "rows": rows, "b_source": "uploaded" if filed else "books",
            "uploaded_by": filed.uploaded_by if filed else None,
            "all_agree": all(r["agrees"] for r in rows),
            "cancelled_in_period": len(g1["cancelled"]),
            "note": ("A is GSTR-1 from the books (outward) and, for input credit, the vendor "
                     "invoices recorded here. B is the GSTR-3B uploaded for the period"
                     + ("" if filed else " — none yet, so B is the GSTR-3B computed from the books")
                     + ". A − B should be nil; a positive figure means liability reported in "
                     "GSTR-1 but not paid through GSTR-3B, a negative one the reverse. Cancelled "
                     "invoices are already out of A; if they were filed in B, upload the revised "
                     "figures or adjust in the next period.")}
